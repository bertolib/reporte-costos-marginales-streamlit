"""Excel-backed market data repository.

The repository reads the workbook as a source of truth without modifying it.
It converts wide Excel sheets into a canonical long table suitable for charts,
KPIs, and a future PostgreSQL implementation.
"""

from pathlib import Path

import openpyxl
import numpy as np
import pandas as pd

from src.domain.enums import Granularity, Scenario
from src.domain.node_catalog import EXCEL_NODE_COLUMNS, NODE_CATALOG
from src.infrastructure.repositories.market_repository import MarketRepository


MONTH_LABELS = {
    1: "ene",
    2: "feb",
    3: "mar",
    4: "abr",
    5: "may",
    6: "jun",
    7: "jul",
    8: "ago",
    9: "sept",
    10: "oct",
    11: "nov",
    12: "dic",
}


def _format_month_label(value: object) -> str:
    timestamp = pd.to_datetime(value, errors="coerce")
    if pd.isna(timestamp):
        return str(value)
    return f"{MONTH_LABELS[timestamp.month]}-{timestamp.strftime('%y')}"


class ExcelMarketRepository(MarketRepository):
    """Read marginal-cost observations from the current weekly Excel file."""

    def __init__(self, source_path: Path) -> None:
        self.source_path = source_path

    def load_observations(self, granularity: Granularity) -> pd.DataFrame:
        """Return observations in canonical long format."""
        if granularity == Granularity.HOURLY:
            return self._load_hourly()
        if granularity == Granularity.FIFTEEN_MINUTES:
            return self._load_fifteen_minutes()
        if granularity == Granularity.DAILY:
            return self._load_daily()
        msg = f"Unsupported granularity: {granularity}"
        raise ValueError(msg)

    def get_last_update(self) -> pd.Timestamp | None:
        """Return latest actual timestamp across hourly observations."""
        df = self.load_observations(Granularity.HOURLY)
        actual = df[df["scenario"] == Scenario.ACTUAL.value]
        if actual.empty:
            return None
        return actual["timestamp"].max()

    def load_daily_market_context(self) -> pd.DataFrame:
        """Return the daily CMg, monthly CMg and commodity-price series."""
        df = pd.read_excel(self.source_path, sheet_name=" CMg Diario")
        df = df.rename(
            columns={
                "Fecha": "date",
                "Cmg Diario": "daily_cmg",
                "Cmg promedio Mensual": "monthly_average_cmg",
                "Programado Cen Mensual": "cen_monthly_projection",
                "Precio Carbon": "coal_price",
                "Precio Petroleo": "oil_price",
            }
        )
        columns = [
            "date",
            "daily_cmg",
            "monthly_average_cmg",
            "cen_monthly_projection",
            "coal_price",
            "oil_price",
        ]
        df = df[columns].copy()
        df["date"] = pd.to_datetime(df["date"], errors="coerce")
        for column in columns[1:]:
            df[column] = pd.to_numeric(df[column], errors="coerce")
        df["cen_monthly_projection"] = df["cen_monthly_projection"].replace(0, np.nan)
        projections = self._load_latest_cen_projection_by_month()
        if projections:
            month = df["date"].dt.to_period("M").dt.to_timestamp()
            mapped_projection = month.map(projections)
            df["cen_monthly_projection"] = df["cen_monthly_projection"].combine_first(mapped_projection)
        return df.dropna(subset=["date"]).sort_values("date")

    def load_cen_projections(self) -> pd.DataFrame:
        """Return the Proyecciones CEN matrix from the daily sheet."""
        header_row, start_col, end_col = self._locate_cen_projection_matrix()
        raw = pd.read_excel(
            self.source_path,
            sheet_name=" CMg Diario",
            header=None,
            skiprows=header_row - 1,
            nrows=20,
            usecols=range(start_col - 1, end_col),
        )
        header = raw.iloc[0]
        data = raw.iloc[1:].dropna(how="all").copy()
        data.columns = ["projection_month", *[_format_month_label(value) for value in header.iloc[1:]]]
        data = data.dropna(subset=["projection_month"])
        data["projection_month"] = data["projection_month"].map(_format_month_label)
        value_columns = [column for column in data.columns if column != "projection_month"]
        data[value_columns] = data[value_columns].replace({np.nan: None})
        return data

    def _load_latest_cen_projection_by_month(self) -> dict[pd.Timestamp, float]:
        header_row, start_col, end_col = self._locate_cen_projection_matrix()
        raw = pd.read_excel(
            self.source_path,
            sheet_name=" CMg Diario",
            header=None,
            skiprows=header_row - 1,
            nrows=2,
            usecols=range(start_col - 1, end_col),
        )
        month_headers = pd.to_datetime(raw.iloc[0, 1:], errors="coerce")
        latest_values = pd.to_numeric(raw.iloc[1, 1:], errors="coerce")
        projections: dict[pd.Timestamp, float] = {}
        for month, value in zip(month_headers, latest_values, strict=False):
            if pd.notna(month) and pd.notna(value):
                projections[pd.Timestamp(month).to_period("M").to_timestamp()] = float(value)
        return projections

    def _locate_cen_projection_matrix(self) -> tuple[int, int, int]:
        """Locate the CEN projection matrix in the daily worksheet.

        Returns the row containing target-month headers and the first/last
        columns of the matrix, using one-based Excel coordinates.
        """
        workbook = openpyxl.load_workbook(self.source_path, read_only=True, data_only=True)
        worksheet = workbook[" CMg Diario"]
        anchor_row: int | None = None
        anchor_col: int | None = None
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.strip().upper() == "PROYECCIONES CEN":
                    anchor_row = cell.row
                    anchor_col = cell.column
                    break
            if anchor_row is not None and anchor_col is not None:
                break

        if anchor_row is None or anchor_col is None:
            return 4584, 12, 37

        header_row = anchor_row + 3
        first_col = max(anchor_col - 1, 1)
        last_col = first_col
        for col in range(first_col, worksheet.max_column + 1):
            if worksheet.cell(header_row, col).value is not None:
                last_col = col
        return header_row, first_col, last_col

    def _load_hourly(self) -> pd.DataFrame:
        df = pd.read_excel(self.source_path, sheet_name="Datos CMg Horario")
        df = df.rename(columns={"Año": "year", "Fecha": "date", "Hora": "hour"})
        return self._wide_nodes_to_long(
            df=df,
            granularity=Granularity.HOURLY,
            date_col="date",
            hour_col="hour",
            minute_col=None,
        )

    def _load_fifteen_minutes(self) -> pd.DataFrame:
        df = pd.read_excel(self.source_path, sheet_name="Datos CMg 15min")
        df = df.rename(columns={"Año": "year", "Fecha": "date", "Hora": "hour", "Minutos": "minute"})
        return self._wide_nodes_to_long(
            df=df,
            granularity=Granularity.FIFTEEN_MINUTES,
            date_col="date",
            hour_col="hour",
            minute_col="minute",
        )

    def _load_daily(self) -> pd.DataFrame:
        df = pd.read_excel(self.source_path, sheet_name=" CMg Diario")
        df = df.rename(
            columns={
                "Fecha": "date",
                "Cmg Diario": "value",
                "Cmg promedio Mensual": "monthly_average",
                "Programado Cen Mensual": "projected_monthly",
            }
        )
        df["timestamp"] = pd.to_datetime(df["date"], errors="coerce")
        df["value"] = pd.to_numeric(df["value"], errors="coerce")
        df = df.dropna(subset=["timestamp", "value"])
        df = df.assign(
            node_id="system_daily",
            node_name="Promedio sistema diario",
            zone="Sistema",
            unit="USD/MWh",
            granularity=Granularity.DAILY.value,
            scenario=Scenario.ACTUAL.value,
            source=self.source_path.name,
        )
        return df[
            [
                "timestamp",
                "node_id",
                "node_name",
                "zone",
                "value",
                "unit",
                "granularity",
                "scenario",
                "source",
            ]
        ].sort_values("timestamp")

    def _wide_nodes_to_long(
        self,
        df: pd.DataFrame,
        granularity: Granularity,
        date_col: str,
        hour_col: str,
        minute_col: str | None,
    ) -> pd.DataFrame:
        available_columns = [col for col in EXCEL_NODE_COLUMNS if col in df.columns]
        base_columns = [date_col, hour_col, *([minute_col] if minute_col else [])]
        df = df[base_columns + available_columns].copy()
        df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
        df[hour_col] = pd.to_numeric(df[hour_col], errors="coerce")
        if minute_col:
            df[minute_col] = self._parse_minute(df[minute_col])
        else:
            df["minute"] = 0
            minute_col = "minute"

        actual_cutoff = self._detect_actual_cutoff(df, date_col, available_columns)
        long_df = df.melt(
            id_vars=[date_col, hour_col, minute_col],
            value_vars=available_columns,
            var_name="excel_node_column",
            value_name="value",
        )
        long_df["value"] = pd.to_numeric(long_df["value"], errors="coerce")
        long_df = long_df.dropna(subset=[date_col, hour_col, "value"])
        long_df["node_id"] = long_df["excel_node_column"].map(EXCEL_NODE_COLUMNS)
        long_df["node_name"] = long_df["node_id"].map(lambda node_id: NODE_CATALOG[node_id].name)
        long_df["zone"] = long_df["node_id"].map(lambda node_id: NODE_CATALOG[node_id].zone)
        long_df["timestamp"] = self._build_timestamp(long_df, date_col, hour_col, minute_col, granularity)
        long_df = long_df.dropna(subset=["timestamp"])
        long_df["scenario"] = Scenario.ACTUAL.value
        if actual_cutoff is not None:
            long_df.loc[long_df["timestamp"].dt.normalize() > actual_cutoff, "scenario"] = Scenario.PROJECTED.value
        long_df = long_df.assign(
            unit="USD/MWh",
            granularity=granularity.value,
            source=self.source_path.name,
        )
        return long_df[
            [
                "timestamp",
                "node_id",
                "node_name",
                "zone",
                "value",
                "unit",
                "granularity",
                "scenario",
                "source",
            ]
        ].sort_values(["timestamp", "node_name"])

    @staticmethod
    def _detect_actual_cutoff(
        df: pd.DataFrame,
        date_col: str,
        value_columns: list[str],
    ) -> pd.Timestamp | None:
        """Detect the latest observed day before the workbook's future gap.

        The workbook keeps future calendar rows for projections/formulas. Actual
        data ends at the last day before the first all-empty day after data has
        started. This avoids treating later formula zeros as observed prices.
        """
        if not value_columns:
            return None

        dates = pd.to_datetime(df[date_col], errors="coerce").dt.normalize()
        values = df[value_columns].apply(pd.to_numeric, errors="coerce")
        by_date = values.notna().groupby(dates).sum().sum(axis=1).sort_index()
        by_date = by_date[by_date.index.notna()]
        if by_date.empty:
            return None

        valid_seen = False
        last_valid_date: pd.Timestamp | None = None
        for date, valid_count in by_date.items():
            if valid_count > 0:
                valid_seen = True
                last_valid_date = pd.Timestamp(date)
                continue
            if valid_seen:
                return last_valid_date
        return last_valid_date

    @staticmethod
    def _parse_minute(series: pd.Series) -> pd.Series:
        parsed = pd.to_timedelta(series.astype(str), errors="coerce")
        return (parsed.dt.components.hours * 60 + parsed.dt.components.minutes).fillna(0)

    @staticmethod
    def _build_timestamp(
        df: pd.DataFrame,
        date_col: str,
        hour_col: str,
        minute_col: str,
        granularity: Granularity,
    ) -> pd.Series:
        hour = pd.to_numeric(df[hour_col], errors="coerce")
        minute = pd.to_numeric(df[minute_col], errors="coerce").fillna(0)
        base_date = pd.to_datetime(df[date_col], errors="coerce")
        if granularity == Granularity.FIFTEEN_MINUTES:
            hour = hour.fillna(0).clip(lower=0)
            return base_date + pd.to_timedelta(hour.astype(int), unit="h") + pd.to_timedelta(minute.astype(int), unit="m")

        hour = hour.fillna(1).clip(lower=1)
        return base_date + pd.to_timedelta(hour.astype(int) - 1, unit="h") + pd.to_timedelta(minute.astype(int), unit="m")
